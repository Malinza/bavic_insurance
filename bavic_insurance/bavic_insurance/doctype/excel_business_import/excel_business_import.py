import frappe
from frappe.model.document import Document
from frappe.utils import flt
from frappe.utils.file_manager import get_file_path
import openpyxl
from datetime import datetime
import time


class ExcelBusinessImport(Document):
	pass


# ---------------------------------------------------------------------------
# Whitelisted API — called from the "Import Now" button
# ---------------------------------------------------------------------------

@frappe.whitelist()
def run_import(docname):
	"""
	Main import function.
	Reads the attached .xlsx, resolves column headers flexibly, ensures all
	masters exist, and creates Insurance Transaction records (optionally submitting).

	Publishes realtime progress events for EVERY input row so the client
	sees a live progress bar and live activity stream.
	"""
	doc = frappe.get_doc("Excel Business Import", docname)

	if doc.status == "In Progress":
		frappe.throw("Import is already in progress for this document.")

	# Reset counters and state
	doc.db_set("status", "In Progress", notify=True)
	doc.db_set("imported_rows", 0)
	doc.db_set("skipped_rows", 0)
	doc.db_set("failed_rows", 0)
	doc.db_set("import_log", "")
	frappe.db.commit()

	try:
		file_path = get_file_path(doc.excel_file)
	except Exception:
		_fail(doc, "❌ Could not resolve attached file path.")
		return {"error": "Could not resolve attached file path."}

	try:
		wb = openpyxl.load_workbook(file_path, data_only=True)
	except Exception as e:
		_fail(doc, f"❌ Could not open workbook: {e}")
		return {"error": str(e)}

	ws = wb.active  # first/only sheet

	# Build normalized header index map
	raw_headers = {}
	norm_headers = {}
	for cell in ws[1]:
		if cell.value is not None:
			raw_val = str(cell.value).strip()
			idx = cell.column - 1  # 0-based index
			raw_headers[raw_val] = idx
			norm_headers[_norm_col(raw_val)] = idx

	def find_col(candidates):
		for cand in candidates:
			if cand in raw_headers:
				return raw_headers[cand]
			cand_norm = _norm_col(cand)
			if cand_norm in norm_headers:
				return norm_headers[cand_norm]
		return None

	# Resolve required columns with flexible aliases
	col_holder    = find_col(["Policy Holder Name", "Policyholder Name", "Policy Holder", "Policyholder", "Customer Name", "Customer"])
	col_eff_date  = find_col(["Effective Date", "Start Date", "Commencement Date"])
	col_ren_date  = find_col(["Renewal Date", "Expiry Date", "End Date"])
	col_post_date = find_col(["PostingDate", "Posting Date", "Date"])
	col_amount    = find_col(["EntryAmount", "Entry Amount", "Amount", "Premium", "Premium Amount"])
	col_product   = find_col(["Product", "Product Name", "Plan", "Insurance Product"])
	col_b_type    = find_col(["Business Type", "Business_Type", "Type of Business"])

	missing = []
	if col_holder is None: missing.append("Policy Holder Name")
	if col_eff_date is None: missing.append("Effective Date")
	if col_ren_date is None: missing.append("Renewal Date")
	if col_post_date is None: missing.append("PostingDate")
	if col_amount is None: missing.append("EntryAmount")
	if col_product is None: missing.append("Product")
	if col_b_type is None: missing.append("Business Type")

	if missing:
		msg = f"❌ Missing required columns: {', '.join(missing)}"
		_fail(doc, msg)
		return {"error": msg}

	# Optional columns with smart fallbacks
	col_agent_name = find_col(["Intermediary Name", "Intermediary", "Agent Name", "Agent"])
	col_agent_type = find_col(["Intermediary Type", "Agent Type"])
	col_journal    = find_col(["Journal Description", "Journal_Description", "Description", "Narration", "Remarks"])
	col_email      = find_col(["PlicyHolder Email", "Policy Holder Email", "Policyholder Email", "Email"])
	col_phone      = find_col(["PlicyHolder Phone", "Policy Holder Phone", "Policyholder Phone", "Phone", "Mobile"])

	# Collect all non-blank data rows
	raw_rows = []
	for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
		if not all(v is None for v in row):
			raw_rows.append((row_num, row))

	total_rows = len(raw_rows)
	if total_rows == 0:
		_fail(doc, "❌ Excel file contains no data rows.")
		return {"error": "Excel file contains no data rows."}

	# Pre-load existing masters in memory to eliminate redundant SQL queries
	known_insurers = set(frappe.get_all("Insurer", pluck="name"))
	known_agents = set(frappe.get_all("Agent", pluck="name"))
	known_agent_types = set(frappe.get_all("Agent Type", pluck="name"))
	known_products = set(frappe.get_all("Insurance Product", pluck="name"))
	product_rates = {
		row.name: flt(row.commission_rate or 0)
		for row in frappe.get_all("Insurance Product", fields=["name", "commission_rate"])
	}
	known_business_types = set(frappe.get_all("Business Type", pluck="name"))
	known_customers = set(frappe.get_all("Customer", pluck="name"))

	# Ensure default insurer "Bavic" exists
	_ensure_insurer("Bavic", known_insurers)

	# Ensure default agent type exists
	_ensure_agent_type("Agent Independent", known_agent_types)

	# Emit init event to client
	_publish_event("excel_import_init", {
		"docname": docname,
		"total": total_rows,
		"status": "In Progress",
	})

	log_lines = [f"🚀 Import started: {total_rows} rows to process."]
	imported = 0
	skipped = 0
	failed = 0
	submit_after = bool(doc.submit_after_import)

	# Adaptive timing based on row count
	sleep_time = 0.002 if total_rows > 300 else 0.012
	commit_interval = 50 if total_rows > 300 else 15

	for idx, (row_num, row) in enumerate(raw_rows):
		current_idx = idx + 1
		percent = round((current_idx / total_rows) * 100, 1)

		def get_val(col_idx):
			if col_idx is not None and col_idx < len(row):
				return row[col_idx]
			return None

		policy_holder_name = _clean(get_val(col_holder))
		effective_date     = _to_date(get_val(col_eff_date))
		renewal_date       = _to_date(get_val(col_ren_date))
		posting_date       = _to_date(get_val(col_post_date))
		amount_val         = get_val(col_amount) or 0
		product_name       = _clean(get_val(col_product))
		business_type_name = _clean(get_val(col_b_type))

		# Optional fields with smart fallbacks
		journal_desc       = _clean(get_val(col_journal)) or ""
		intermediary_name  = _clean(get_val(col_agent_name)) or "BAVIC INSURANCE AGENCY"
		intermediary_type  = _clean(get_val(col_agent_type)) or "Agent Independent"
		email              = _clean(get_val(col_email))
		phone              = _clean(get_val(col_phone))

		try:
			amount = float(amount_val)
		except (ValueError, TypeError):
			amount = 0.0

		# Check mandatory fields
		if not all([policy_holder_name, effective_date, renewal_date, posting_date, product_name, business_type_name]):
			failed += 1
			row_status = "failed"
			detail = "Missing required column(s)"
			log_line = f"⚠️ Row {row_num}: Missing mandatory field(s) — {policy_holder_name or 'Unknown'}"
			log_lines.append(log_line)

			_publish_event("excel_import_row", {
				"docname": docname,
				"row_num": row_num,
				"current": current_idx,
				"total": total_rows,
				"percent": percent,
				"status": row_status,
				"customer": policy_holder_name or "Missing Name",
				"product": product_name or "-",
				"business_type": business_type_name or "-",
				"amount": amount,
				"effective_date": str(effective_date) if effective_date else "-",
				"detail": detail,
				"imported": imported,
				"skipped": skipped,
				"failed": failed,
			})
			continue

		try:
			# 1 — Masters get-or-create (in-memory cached)
			_ensure_agent_type(intermediary_type, known_agent_types)
			_ensure_agent(intermediary_name, intermediary_type, known_agents)
			commission_rate = _ensure_product(product_name, known_products, product_rates)
			_ensure_business_type(business_type_name, known_business_types)
			customer_name = _ensure_customer(policy_holder_name, email, phone, known_customers)

			# 2 — Duplicate check
			if _is_duplicate(customer_name, effective_date, amount, product_name, journal_desc):
				skipped += 1
				row_status = "skipped"
				detail = f"Duplicate record — already exists in Insurance Transactions"
				log_line = f"⏭️ Row {row_num}: Skipped duplicate — {policy_holder_name} | {product_name} | {effective_date}"
				log_lines.append(log_line)

				_publish_event("excel_import_row", {
					"docname": docname,
					"row_num": row_num,
					"current": current_idx,
					"total": total_rows,
					"percent": percent,
					"status": row_status,
					"customer": policy_holder_name,
					"product": product_name,
					"business_type": business_type_name,
					"amount": amount,
					"effective_date": str(effective_date),
					"detail": detail,
					"imported": imported,
					"skipped": skipped,
					"failed": failed,
				})
				continue

			# 3 — Create Insurance Transaction
			txn = frappe.get_doc({
				"doctype":              "Insurance Transaction",
				"naming_series":        "INS-.YYYY.-.",
				"posting_date":         posting_date,
				"customer":             customer_name,
				"policy_holder_name":   policy_holder_name,
				"policy_holder_email":  email or "",
				"policy_holder_phone":  phone or "",
				"insurer":              "Bavic",
				"product":              product_name,
				"business_type":        business_type_name,
				"intermediary":         intermediary_name,
				"effective_date":       effective_date,
				"renewal_date":         renewal_date,
				"amount":               amount,
				"commission_rate":      commission_rate,
				"premium_payment":      "Full Premium",
				"journal_description":  journal_desc,
			})
			txn.insert(ignore_permissions=True)

			if submit_after:
				txn.submit()
				row_status = "submitted"
				detail = f"Created & Submitted: {txn.name}"
			else:
				row_status = "imported"
				detail = f"Created Draft: {txn.name}"

			imported += 1
			log_line = f"✅ Row {row_num}: {'Submitted' if submit_after else 'Imported'} — {policy_holder_name} | {product_name} | {amount:,.2f} TZS"
			log_lines.append(log_line)

			_publish_event("excel_import_row", {
				"docname": docname,
				"row_num": row_num,
				"current": current_idx,
				"total": total_rows,
				"percent": percent,
				"status": row_status,
				"customer": policy_holder_name,
				"product": product_name,
				"business_type": business_type_name,
				"amount": amount,
				"effective_date": str(effective_date),
				"detail": detail,
				"imported": imported,
				"skipped": skipped,
				"failed": failed,
			})

		except Exception as e:
			frappe.log_error(frappe.get_traceback(), f"Excel Import Row {row_num}")
			failed += 1
			row_status = "failed"
			detail = f"Error: {str(e)}"
			log_line = f"❌ Row {row_num}: Error — {policy_holder_name} | {e}"
			log_lines.append(log_line)

			_publish_event("excel_import_row", {
				"docname": docname,
				"row_num": row_num,
				"current": current_idx,
				"total": total_rows,
				"percent": percent,
				"status": row_status,
				"customer": policy_holder_name or "Unknown",
				"product": product_name or "-",
				"business_type": business_type_name or "-",
				"amount": amount,
				"effective_date": str(effective_date) if effective_date else "-",
				"detail": detail,
				"imported": imported,
				"skipped": skipped,
				"failed": failed,
			})

		# Smooth WebSocket event streaming pacing
		if sleep_time > 0:
			time.sleep(sleep_time)

		# Periodically commit database so records are saved
		if current_idx % commit_interval == 0:
			frappe.db.commit()

	frappe.db.commit()

	# Finalize document state
	final_status = "Completed" if failed == 0 else "Completed with Errors"
	summary = f"\n🏁 Import Complete — {imported} imported, {skipped} skipped (duplicates), {failed} failed."
	log_lines.append(summary)

	doc.db_set("imported_rows", imported)
	doc.db_set("skipped_rows", skipped)
	doc.db_set("failed_rows", failed)
	doc.db_set("status", final_status)
	doc.db_set("import_log", "\n".join(log_lines))
	frappe.db.commit()

	_publish_event("excel_import_complete", {
		"docname": docname,
		"total": total_rows,
		"imported": imported,
		"skipped": skipped,
		"failed": failed,
		"status": final_status,
	})

	return {
		"imported": imported,
		"skipped": skipped,
		"failed": failed,
		"status": final_status,
	}


# ---------------------------------------------------------------------------
# Master Helpers with in-memory caching
# ---------------------------------------------------------------------------

def _ensure_insurer(name, cache=None):
	if cache is not None and name in cache:
		return
	if not frappe.db.exists("Insurer", name):
		frappe.get_doc({"doctype": "Insurer", "insurer_name": name}).insert(ignore_permissions=True)
		frappe.db.commit()
	if cache is not None:
		cache.add(name)


def _ensure_agent_type(type_name, cache=None):
	if not type_name:
		type_name = "Agent Independent"
	if cache is not None and type_name in cache:
		return
	if not frappe.db.exists("Agent Type", type_name):
		frappe.get_doc({"doctype": "Agent Type", "type_name": type_name}).insert(ignore_permissions=True)
		frappe.db.commit()
	if cache is not None:
		cache.add(type_name)


def _ensure_agent(intermediary_name, intermediary_type=None, cache=None):
	if not intermediary_name:
		intermediary_name = "BAVIC INSURANCE AGENCY"
	if cache is not None and intermediary_name in cache:
		return intermediary_name
	if not frappe.db.exists("Agent", intermediary_name):
		type_to_use = intermediary_type or "Agent Independent"
		_ensure_agent_type(type_to_use)
		frappe.get_doc({
			"doctype": "Agent",
			"intermediary_name": intermediary_name,
			"intermediary_type": type_to_use,
		}).insert(ignore_permissions=True)
		frappe.db.commit()
	if cache is not None:
		cache.add(intermediary_name)
	return intermediary_name


def _ensure_product(product_name, cache=None, rate_cache=None):
	"""Create the product when missing, and return its commission rate.

	The import file does not carry a rate. A product that already exists keeps
	the rate stored on Insurance Product. A product created here starts at 0.
	"""
	if not product_name:
		return 0.0
	if rate_cache is not None and product_name in rate_cache:
		if cache is not None:
			cache.add(product_name)
		return flt(rate_cache[product_name] or 0)

	if frappe.db.exists("Insurance Product", product_name):
		rate = flt(frappe.db.get_value("Insurance Product", product_name, "commission_rate") or 0)
	else:
		frappe.get_doc({
			"doctype": "Insurance Product",
			"product_name": product_name,
			"commission_rate": 0,
		}).insert(ignore_permissions=True)
		frappe.db.commit()
		rate = 0.0

	if cache is not None:
		cache.add(product_name)
	if rate_cache is not None:
		rate_cache[product_name] = rate
	return rate


def _ensure_business_type(business_type_name, cache=None):
	if not business_type_name:
		return
	if cache is not None and business_type_name in cache:
		return
	if not frappe.db.exists("Business Type", business_type_name):
		frappe.get_doc({"doctype": "Business Type", "business_type_name": business_type_name}).insert(ignore_permissions=True)
		frappe.db.commit()
	if cache is not None:
		cache.add(business_type_name)


def _ensure_customer(policy_holder_name, email, phone, cache=None):
	if cache is not None and policy_holder_name in cache:
		return policy_holder_name

	if not frappe.db.exists("Customer", policy_holder_name):
		customer = frappe.get_doc({
			"doctype": "Customer",
			"customer_name": policy_holder_name,
			"policy_holders": [{
				"policy_holder_name": policy_holder_name,
				"email": email or "",
				"phone": phone or "",
			}],
		})
		customer.insert(ignore_permissions=True)
		frappe.db.commit()
	else:
		customer = frappe.get_doc("Customer", policy_holder_name)
		existing = [ph.policy_holder_name for ph in customer.get("policy_holders", [])]
		if policy_holder_name not in existing:
			customer.append("policy_holders", {
				"policy_holder_name": policy_holder_name,
				"email": email or "",
				"phone": phone or "",
			})
			customer.save(ignore_permissions=True)
			frappe.db.commit()

	if cache is not None:
		cache.add(policy_holder_name)
	return policy_holder_name


# ---------------------------------------------------------------------------
# Duplicate Check
# ---------------------------------------------------------------------------

def _is_duplicate(customer, effective_date, amount, product, journal_description=""):
	filters = {
		"customer":            customer,
		"effective_date":      effective_date,
		"amount":              float(amount),
		"product":             product,
	}
	if journal_description:
		filters["journal_description"] = journal_description
	return bool(frappe.db.exists("Insurance Transaction", filters))


# ---------------------------------------------------------------------------
# Utility & Realtime Helpers
# ---------------------------------------------------------------------------

def _norm_col(col_name):
	"""Normalize column header for robust fuzzy matching."""
	if not col_name:
		return ""
	return str(col_name).strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def _clean(value):
	if value is None:
		return None
	s = str(value).strip()
	return s if s else None


def _to_date(value):
	if value is None:
		return None
	if isinstance(value, datetime):
		return value.date().isoformat()
	if hasattr(value, "isoformat"):
		return value.isoformat()
	for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
		try:
			return datetime.strptime(str(value).strip(), fmt).date().isoformat()
		except ValueError:
			continue
	return None


def _publish_event(event_name, payload):
	"""
	Publish realtime event to both the user's private room and Desk site room.
	Ensures immediate client delivery without waiting for database commits.
	"""
	if frappe.session and frappe.session.user:
		frappe.publish_realtime(
			event=event_name,
			message=payload,
			user=frappe.session.user,
			after_commit=False,
		)

	frappe.publish_realtime(
		event=event_name,
		message=payload,
		after_commit=False,
	)


def _fail(doc, message):
	doc.db_set("status", "Failed")
	doc.db_set("import_log", message)
	frappe.db.commit()
	_publish_event("excel_import_complete", {
		"docname": doc.name,
		"status": "Failed",
		"error": message,
	})
